from pathlib import Path
import yaml
from openpyxl import load_workbook

input_xlsx = Path("scenario_data/metal_primary_secondary_shares.xlsx")
output_yaml = Path("scenario_data/metal_primary_secondary_shares.yaml")

wb = load_workbook(input_xlsx, data_only=False)
ws = wb["Shares"]

result = {}
row = 5
while True:
    key = ws[f"A{row}"].value
    if key in (None, ""):
        break

    result[str(key)] = {
        "name": ws[f"B{row}"].value,
        "reference product": ws[f"C{row}"].value,
        "shares": {
            "primary": {
                2020: float(ws[f"D{row}"].value),
                2050: float(ws[f"E{row}"].value),
            },
            "secondary": {
                2020: float(ws[f"F{row}"].value),
                2050: float(ws[f"G{row}"].value),
            },
        },
    }
    row += 1

with open(output_yaml, "w", encoding="utf-8") as f:
    yaml.safe_dump(result, f, sort_keys=False, allow_unicode=True, default_flow_style=False)

print(f"Wrote {output_yaml}")