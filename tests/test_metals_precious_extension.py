"""Reviewed precious-metal extension: real policy files, synthetic inventories.

These tests do not build Brightway databases or claim to validate a future
scrap-stock forecast. They exercise the configured policy and supplier boundary.
"""
from copy import deepcopy
import json
from pathlib import Path
import unittest
from unittest.mock import patch

from metals_scenario_policy import (
    MetalPolicyError, MetalSteelPolicy, apply_database, classify_supply,
    plan_database,
)
from tests.test_metals_scenario_policy import configuration, exchange, inventory, market
from primary_vs_secondary_metals.share_files import (
    GENERATOR_VERSION, ShareDataConflictError, ShareDataError, _generated_cases, _inputs,
    load_yaml,
)


ROOT = Path(__file__).resolve().parents[1]
POLICY = ROOT / "configuration_file" / "metals_scenario_policy.yaml"
ELECTRONICS = (
    "treatment of precious metal from electronics scrap, in anode slime, "
    "precious metal extraction"
)
CATALYST = "treatment of automobile catalyst"
PARKES = "treatment of crust from Parkes process for lead production"
X_RAY = "treatment of waste x-ray film"
OBSERVED_2025 = {
    "gold": 1404.3 / (3671.6 + 1404.3),
    "silver": 197.6 / (846.6 + 197.6),
    "platinum": 1547 / (5557 + 1547),
    "palladium": 3135 / (6589 + 3135),
}
CATALOGUE_2020 = {
    "gold": 0.004426657,
    "silver": 0.084636523,
    "platinum": 0.044600386,
    "palladium": 0.155485524,
}
CASE_INCREMENT = {
    "conservative": 0.0,
    "moderate_recycling": 0.10,
    "high_recycling": 0.20,
}


def precious_market(product, *, secondary_share=0.01, location="World", gross=1.0):
    """Include multiple routes, distribution losses and untouched services."""
    name = "market for " + product
    if product == "platinum":
        recycled = ((CATALYST, "RER", 0.4), (CATALYST, "RoW", 0.6))
    elif product == "palladium":
        recycled = ((CATALYST, "RER", 0.3), (ELECTRONICS, "SE", 0.7))
    elif product == "silver":
        recycled = ((ELECTRONICS, "SE", 0.4), (X_RAY, "GLO", 0.6))
    else:
        # The native metadata includes SE although the upstream gold selector
        # uses GLO. Local exact-name matching must not lose the real SE route.
        recycled = ((ELECTRONICS, "SE", 0.4), (ELECTRONICS, "RoW", 0.6))
    remaining = gross * (1 - secondary_share)
    supplies = [exchange(product + " production, primary", remaining * 0.7,
                         product=product, location="RoW")]
    if product == "silver":
        supplies.append(exchange(PARKES, remaining * 0.3, product=product, location="GLO",
                                 **{"uncertainty type": 3, "loc": remaining * 0.3,
                                    "scale": 0.005, "minimum": remaining * 0.2,
                                    "maximum": remaining * 0.4,
                                    "comment": "Primary-metal co-product, not recovered scrap"}))
    else:
        supplies.append(exchange(product + " production, primary, other route", remaining * 0.3,
                                 product=product, location="GLO"))
    supplies.extend(exchange(provider, gross * secondary_share * weight,
                             product=product, location=region,
                             **{"comment": "Inherited burden-carrying recipe"})
                    for provider, region, weight in recycled)
    services = [exchange("transport, freight", 0.02, product="transport", unit="ton kilometer"),
                {"name": "Carbon dioxide, fossil", "type": "biosphere", "amount": 0.04}]
    return {"name": name, "reference product": product, "unit": "kilogram", "location": location,
            "exchanges": [{"name": name, "product": product, "unit": "kilogram", "amount": 1,
                           "location": location, "type": "production"}] + supplies + services}


class PreciousRepositoryPolicyTests(unittest.TestCase):
    def policy(self, case="baseline"):
        return MetalSteelPolicy.from_file(POLICY, nonferrous_case=case)

    @staticmethod
    def row_for(plan, product):
        return next(row for row in plan.audit if row["reference product"] == product)

    def test_eleven_market_products_are_explicitly_configured_and_native_baseline_is_noop(self):
        selected = self.policy()
        self.assertEqual(len(selected.rules), 11)
        self.assertEqual({rule["reference product"] for rule in selected.rules},
                         {"aluminium, cast alloy", "copper, cathode", "nickel, class 1",
                          "cobalt", "lead", "lithium carbonate, battery grade", "lithium carbonate", *OBSERVED_2025})
        self.assertEqual(set(selected.source_hashes), {str(POLICY.resolve())})
        db = sum((inventory(precious_market(product)) for product in OBSERVED_2025), [])
        # Even an anomalous native steel recipe must be preserved exactly.
        db.append(market(steel=True))
        before = deepcopy(db)
        plan = plan_database(db, selected, scenario={"year": 2050})
        self.assertEqual(plan.changes, ())
        self.assertEqual(apply_database(db, plan, policy=selected), before)
        self.assertTrue(plan.stats["native_baseline"])
        self.assertEqual(plan.stats["steel_owner"], "endogenous IAM/EDM-I")

    def test_current_anchors_and_future_sensitivity_offsets_are_not_confused(self):
        for case, increment in CASE_INCREMENT.items():
            selected = self.policy(case)
            for product, anchor in OBSERVED_2025.items():
                with self.subTest(case=case, product=product):
                    rule = next(rule for rule in selected.rules if rule["reference product"] == product)
                    curve = rule["targets"][case]
                    self.assertEqual(set(curve), {2020, 2025, 2050})
                    self.assertAlmostEqual(curve[2020], CATALOGUE_2020[product])
                    self.assertAlmostEqual(curve[2025], anchor)
                    self.assertAlmostEqual(curve[2050], anchor + increment)
                    self.assertEqual(set(rule["locations"]), {"World", "GLO"})

    def test_linear_2025_2050_floors_preserve_gross_supply_recipe_and_services(self):
        for case, increment in CASE_INCREMENT.items():
            selected = self.policy(case)
            for product, anchor in OBSERVED_2025.items():
                for year in (2025, 2030, 2035, 2040, 2045, 2050):
                    with self.subTest(case=case, product=product, year=year):
                        ds = precious_market(product, gross=1.01)
                        db = inventory(ds)
                        plan = plan_database(db, selected, scenario={"year": year})
                        expected = anchor + increment * (year - 2025) / 25
                        audit = self.row_for(plan, product)
                        self.assertAlmostEqual(audit["applied_secondary_share"], expected)
                        revised = apply_database(db, plan, policy=selected)[0]
                        # Exact same suppliers, amounts alone rescaled within each
                        # supply role, and native distribution losses retained.
                        self.assertEqual([e.get("name") for e in revised["exchanges"]],
                                         [e.get("name") for e in ds["exchanges"]])
                        inputs = revised["exchanges"][1:5]
                        self.assertAlmostEqual(sum(e["amount"] for e in inputs), 1.01)
                        self.assertAlmostEqual(sum(e["amount"] for e in inputs[2:]), 1.01 * expected)
                        self.assertAlmostEqual(inputs[0]["amount"] / inputs[1]["amount"], 7 / 3)
                        if product == "palladium":
                            self.assertAlmostEqual(inputs[2]["amount"] / inputs[3]["amount"], 3 / 7)
                        else:
                            self.assertAlmostEqual(inputs[2]["amount"] / inputs[3]["amount"], 2 / 3)
                        self.assertEqual(revised["exchanges"][5:], ds["exchanges"][5:])
                        self.assertEqual(db[0], ds)  # Planning/applying never mutates the input.

    def test_new_curves_clamp_and_retain_higher_native_shares(self):
        selected = self.policy("high_recycling")
        for product, anchor in OBSERVED_2025.items():
            for year, floor in ((2010, CATALOGUE_2020[product]),
                                (2020, CATALOGUE_2020[product]),
                                (2025, anchor), (2050, anchor + 0.2), (2060, anchor + 0.2)):
                with self.subTest(product=product, year=year):
                    ds = precious_market(product, secondary_share=0.0)
                    # Zero rows exist but are not positive recipes. A source
                    # inventory with identical verified providers is supplied.
                    source = precious_market(product, secondary_share=0.01)
                    plan = plan_database(inventory(ds), selected, scenario={"year": year},
                                         source_database=[source])
                    self.assertAlmostEqual(self.row_for(plan, product)["applied_secondary_share"], floor)
            db = inventory(precious_market(product, secondary_share=0.9))
            before = deepcopy(db)
            plan = plan_database(db, selected, scenario={"year": 2030})
            self.assertEqual(plan.changes, ())
            self.assertAlmostEqual(self.row_for(plan, product)["applied_secondary_share"], 0.9)
            self.assertEqual(apply_database(db, plan, policy=selected), before)

    def test_reference_stays_original_2020_and_primary_only_is_zero(self):
        for case in ("reference", "primary_only"):
            selected = self.policy(case)
            for product in OBSERVED_2025:
                with self.subTest(case=case, product=product):
                    ds = precious_market(product, secondary_share=0.2)
                    plan = plan_database(inventory(ds), selected, scenario={"year": 2040})
                    expected = CATALOGUE_2020[product] if case == "reference" else 0.0
                    self.assertAlmostEqual(self.row_for(plan, product)["applied_secondary_share"], expected)

    def test_original_six_2050_endpoints_follow_revised_evidence(self):
        expected = {
            "aluminium, cast alloy": (0.45, 0.55, 0.65),
            "copper, cathode": (0.25, 0.35, 0.45),
            "nickel, class 1": (0.15, 0.22, 0.30),
            "cobalt": (0.20, 0.30, 0.40),
            "lead": (0.65, 0.77, 0.85),
            "lithium carbonate, battery grade": (0.15, 250 / 1122, 560 / 1501),
        }
        for index, case in enumerate(CASE_INCREMENT):
            selected = self.policy(case)
            for rule in selected.rules:
                product = rule["reference product"]
                if product in expected:
                    with self.subTest(case=case, product=product):
                        self.assertAlmostEqual(rule["targets"][case][2050], expected[product][index])

    def test_new_candidate_evidence_does_not_activate_unsupported_runtime_markets(self):
        selected = self.policy("high_recycling")
        products = ("dysprosium oxide", "terbium oxide", "praseodymium oxide", "neodymium oxide",
                    "cerium oxide", "lanthanum oxide", "yttrium oxide", "graphite", "zinc",
                    "lithium hydroxide, battery grade", "steel, low-alloyed")
        configured = {rule["reference product"] for rule in selected.rules}
        self.assertFalse(configured.intersection(products))
        db = []
        for product in products:
            name = "market for " + product
            db.append({"name": name, "reference product": product, "unit": "kilogram", "location": "GLO",
                       "exchanges": [{"name": name, "product": product, "unit": "kilogram",
                                      "location": "GLO", "amount": 1, "type": "production"}]})
        before = deepcopy(db)
        plan = plan_database(db, selected, scenario={"year": 2050})
        self.assertEqual(plan.changes, ())
        self.assertEqual(plan.stats["matched_markets"], 0)
        self.assertEqual(apply_database(db, plan, policy=selected), before)

    def test_silver_parkes_is_explicit_primary_remainder_not_secondary(self):
        selected = self.policy("high_recycling")
        rule = next(rule for rule in selected.rules if rule["reference product"] == "silver")
        ds = precious_market("silver")
        parkes = ds["exchanges"][2]
        self.assertEqual(classify_supply(parkes, rule), "primary")
        self.assertEqual(classify_supply(ds["exchanges"][3], rule), "secondary")
        self.assertEqual(classify_supply(ds["exchanges"][4], rule), "secondary")
        plan = plan_database(inventory(ds), selected, scenario={"year": 2050})
        revised = apply_database(inventory(ds), plan, policy=selected)[0]["exchanges"][2]
        factor = revised["amount"] / parkes["amount"]
        self.assertEqual(revised["uncertainty type"], parkes["uncertainty type"])
        for field in ("loc", "scale", "minimum", "maximum"):
            self.assertAlmostEqual(revised[field], parkes[field] * factor)
        self.assertEqual(revised["comment"], parkes["comment"])
        # The exception is exact, not a broad licence for other treatment names.
        self.assertEqual(classify_supply(dict(parkes, name=PARKES + ", unreviewed"), rule), "unknown")
        self.assertIsNone(classify_supply(dict(parkes, product="lead"), rule))
        self.assertIsNone(classify_supply(dict(parkes, amount=-0.01), rule))

    def test_gold_se_and_only_reviewed_local_recovery_names_are_accepted(self):
        selected = self.policy("high_recycling")
        for product in ("gold", "silver"):
            rule = next(rule for rule in selected.rules if rule["reference product"] == product)
            with self.subTest(product=product):
                self.assertTrue(rule["strict_secondary_matching"])
                known = exchange(ELECTRONICS, 0.1, product=product, location="SE")
                self.assertEqual(classify_supply(known, rule), "secondary")
                unknown = dict(known, name="treatment of unreviewed precious-metal waste")
                self.assertEqual(classify_supply(unknown, rule), "unknown")
                ds = precious_market(product)
                ds["exchanges"][3] = unknown
                with self.assertRaisesRegex(MetalPolicyError, "Unrecognised"):
                    plan_database(inventory(ds), selected, scenario={"year": 2050})

    def test_reports_carry_scientific_assumptions_and_applied_supplier_evidence(self):
        selected = self.policy("high_recycling")
        for product in OBSERVED_2025:
            with self.subTest(product=product):
                ds = precious_market(product)
                plan = plan_database(inventory(ds), selected, scenario={"year": 2050})
                row = self.row_for(plan, product)
                self.assertTrue(row["evidence_ids"])
                self.assertTrue(row["evidence_details"])
                self.assertTrue(row["target_note"])
                self.assertTrue(row["secondary_matching_note"])
                self.assertTrue(row["applied_primary_suppliers"])
                self.assertTrue(row["applied_secondary_suppliers"])
                self.assertEqual(row["share_data_targets"][2025], OBSERVED_2025[product])
                self.assertEqual(len(row["share_data_sha256"]), 64)
                if product == "silver":
                    self.assertIn(PARKES, row["primary_remainder_names"])
                    self.assertTrue(row["primary_remainder_note"])


class PrimaryRemainderSafeguardTests(unittest.TestCase):
    def config(self):
        config = configuration()
        config["markets"][0].update(
            primary_remainder_names=["treatment of primary co-product"],
            primary_remainder_note="Reviewed primary-derived intermediate, not recovered scrap.",
            strict_secondary_matching=True,
            secondary={"names": ["treatment of copper scrap, refining"]},
        )
        return config

    def test_primary_remainder_requires_note_exact_list_and_no_secondary_overlap(self):
        variants = []
        missing_note = self.config()
        missing_note["markets"][0].pop("primary_remainder_note")
        variants.append(missing_note)
        blank_note = self.config()
        blank_note["markets"][0]["primary_remainder_note"] = " "
        variants.append(blank_note)
        string_names = self.config()
        string_names["markets"][0]["primary_remainder_names"] = "treatment of primary co-product"
        variants.append(string_names)
        blank_name = self.config()
        blank_name["markets"][0]["primary_remainder_names"] = [""]
        variants.append(blank_name)
        overlap = self.config()
        overlap["markets"][0]["primary_remainder_names"] = ["treatment of copper scrap, refining"]
        variants.append(overlap)
        for config in variants:
            with self.subTest(config=config):
                with self.assertRaises(MetalPolicyError):
                    MetalSteelPolicy.from_mapping(config, nonferrous_case="high_recycling")

    def test_strict_matching_must_be_boolean(self):
        for invalid in ("yes", 1, None):
            config = self.config()
            config["markets"][0]["strict_secondary_matching"] = invalid
            with self.subTest(invalid=invalid):
                with self.assertRaises(MetalPolicyError):
                    MetalSteelPolicy.from_mapping(config, nonferrous_case="high_recycling")

    def test_unknown_treatment_stays_unknown_and_aggregate_stays_aggregate(self):
        selected = MetalSteelPolicy.from_mapping(self.config(), nonferrous_case="high_recycling")
        rule = selected.rules[0]
        self.assertEqual(classify_supply(exchange("treatment of primary co-product", 0.1), rule), "primary")
        self.assertEqual(classify_supply(exchange("treatment of primary co-product, other", 0.1), rule), "unknown")
        self.assertEqual(classify_supply(exchange("market for copper, cathode", 1), rule), "aggregate")
        self.assertEqual(classify_supply(exchange("copper production, primary", 0.8), rule), "primary")


class EvidenceAnchorValidationTests(unittest.TestCase):
    """CSV anchor validation without modifying repository inputs or outputs."""

    @staticmethod
    def baseline():
        return {"gold": {"name": "market for gold", "reference product": "gold",
                         "shares": {"primary": {2020: 0.99, 2050: 0.99},
                                    "secondary": {2020: 0.01, 2050: 0.01}}}}

    @staticmethod
    def record(**extra):
        return {"family": "gold", "secondary_share_2050_conservative": "0.3",
                "secondary_share_2050_baseline": "0.4",
                "secondary_share_2050_optimistic": "0.5",
                "basis_type": "observed anchor with analyst sensitivity",
                "confidence": "qualified", "source_short": "fixture source",
                "source_url": "https://example.org/source", "notes": "Fixture only.", **extra}

    def inputs(self, row, baseline=None):
        baseline = self.baseline() if baseline is None else baseline
        with patch("primary_vs_secondary_metals.share_files._read_csv", side_effect=[
                [row], [{"product": "gold", "family": "gold"}]]):
            return _inputs(baseline, "fixture_evidence.csv", "fixture_mapping.csv")

    def test_anchor_columns_are_optional_but_a_nonblank_pair_is_required(self):
        for extra in ({}, {"anchor_year": "", "anchor_secondary_share": ""}):
            evidence, _ = self.inputs(self.record(**extra))
            self.assertNotIn("anchor_year", evidence["gold"])
            self.assertNotIn("anchor_secondary_share", evidence["gold"])
        for extra in ({"anchor_year": "2025"}, {"anchor_secondary_share": "0.25"},
                      {"anchor_year": "2025", "anchor_secondary_share": ""},
                      {"anchor_year": "", "anchor_secondary_share": "0"}):
            with self.subTest(extra=extra):
                with self.assertRaisesRegex(ShareDataError, "supplied together"):
                    self.inputs(self.record(**extra))

    def test_anchor_year_and_share_reject_invalid_nonfinite_and_out_of_bounds_values(self):
        for year in ("2020", "2050", "2025.5", "nan", "inf", "true", "future"):
            with self.subTest(year=year):
                with self.assertRaisesRegex(ShareDataError, "anchor_year"):
                    self.inputs(self.record(anchor_year=year, anchor_secondary_share="0.25"))
        for share in ("nan", "inf", "-inf", "-0.01", "1.01", "not a share"):
            with self.subTest(share=share):
                with self.assertRaises(ShareDataError):
                    self.inputs(self.record(anchor_year="2025", anchor_secondary_share=share))

    def test_explicit_zero_anchor_is_retained_and_original_source_is_untouched(self):
        baseline = self.baseline()
        before = deepcopy(baseline)
        evidence, mapping = self.inputs(self.record(anchor_year="2025", anchor_secondary_share="0"),
                                        baseline=baseline)
        self.assertEqual(evidence["gold"]["anchor_year"], 2025)
        self.assertEqual(evidence["gold"]["anchor_secondary_share"], 0.0)
        generated = _generated_cases(baseline, evidence, mapping)
        self.assertEqual(baseline, before)
        for case in ("conservative", "baseline", "optimistic"):
            with self.subTest(case=case):
                shares = generated[case]["gold"]["shares"]
                self.assertEqual(shares["secondary"][2025], 0)
                self.assertEqual(shares["primary"][2025], 1)
                self.assertEqual(shares["secondary"][2020], 0.01)
                self.assertEqual(shares["primary"][2020], 0.99)

    def test_reference_only_scope_skips_both_anchor_and_endpoint(self):
        baseline = self.baseline()
        before = deepcopy(baseline)
        with patch("primary_vs_secondary_metals.share_files._read_csv", side_effect=[
                [self.record(anchor_year="2025", anchor_secondary_share="0.25")],
                [{"product": "gold", "family": "gold", "evidence_scope": "reference_only"}]]):
            evidence, mapping = _inputs(baseline, "fixture_evidence.csv", "fixture_mapping.csv")
        generated = _generated_cases(baseline, evidence, mapping)
        for case in ("conservative", "baseline", "optimistic"):
            self.assertEqual(generated[case], before)
        self.assertEqual(baseline, before)
        self.assertEqual(evidence["gold"]["excluded_products"], ["gold"])

    def test_real_rare_earth_scope_holds_thirteen_oxides_and_only_generates_four_magnet_candidates(self):
        directory = ROOT / "primary_vs_secondary_metals/scenario_data"
        baseline = load_yaml(directory / "metal_primary_secondary_shares.yaml")
        before = deepcopy(baseline)
        evidence, mapping = _inputs(baseline, directory / "metal_2050_literature_evidence.csv",
                                    directory / "metal_product_family_mapping.csv")
        all_rare_earths = {product for product, family in mapping.items() if family == "rare_earths"}
        allowed = {"dysprosium oxide", "terbium oxide", "praseodymium oxide", "neodymium oxide"}
        self.assertEqual(len(all_rare_earths), 17)
        self.assertEqual(set(evidence["rare_earths"]["excluded_products"]), all_rare_earths - allowed)
        generated = _generated_cases(baseline, evidence, mapping)
        for case, endpoint in (("conservative", 0.20), ("baseline", 0.30), ("optimistic", 0.39)):
            for product in all_rare_earths:
                with self.subTest(case=case, product=product):
                    if product in allowed:
                        self.assertAlmostEqual(generated[case][product]["shares"]["secondary"][2050], endpoint)
                        self.assertEqual(generated[case][product]["shares"]["secondary"][2020],
                                         before[product]["shares"]["secondary"][2020])
                    else:
                        self.assertEqual(generated[case][product], before[product])
        self.assertEqual(baseline, before)

    def test_conflicting_existing_anchor_raises_without_overwriting_reference_curves(self):
        baseline = self.baseline()
        baseline["gold"]["shares"]["secondary"][2025] = 0.2
        baseline["gold"]["shares"]["primary"][2025] = 0.8
        before = deepcopy(baseline)
        evidence, mapping = self.inputs(self.record(anchor_year="2025", anchor_secondary_share="0.25"),
                                        baseline=baseline)
        with self.assertRaisesRegex(ShareDataConflictError, "source curve conflicts"):
            _generated_cases(baseline, evidence, mapping)
        self.assertEqual(baseline, before)
        # An explicitly matching anchor is accepted, without mutating any source
        # year. The 2050 target belongs only to the generated scenario copies.
        evidence["gold"]["anchor_secondary_share"] = 0.2
        generated = _generated_cases(baseline, evidence, mapping)
        self.assertEqual(baseline, before)
        self.assertEqual(generated["optimistic"]["gold"]["shares"]["secondary"][2025], 0.2)
        self.assertEqual(generated["optimistic"]["gold"]["shares"]["secondary"][2050], 0.5)

    def test_real_generation_receipt_has_scoped_version_three_and_no_manual_outputs(self):
        manifest = json.loads((ROOT / "primary_vs_secondary_metals/scenario_data/"
                               "share_generation_manifest.json").read_text(encoding="utf-8"))
        self.assertEqual(GENERATOR_VERSION, 3)
        self.assertEqual(manifest["generator_version"], GENERATOR_VERSION)
        self.assertEqual(manifest["manual_outputs"], [])
        self.assertEqual(manifest["product_count"], 104)


class MetalsWorkbookConsistencyTests(unittest.TestCase):
    """Read-only checks prevent documentation drifting away from live shares."""

    def setUp(self):
        try:
            from openpyxl import load_workbook
        except ImportError:
            self.skipTest("Optional workbook inspection dependency openpyxl is unavailable")
        self.path = ROOT / "primary_vs_secondary_metals/scenario_data/metal_primary_secondary_shares.xlsx"
        self.values = load_workbook(self.path, read_only=True, data_only=True)
        self.addCleanup(self.values.close)
        self.formulas = load_workbook(self.path, read_only=True, data_only=False)
        self.addCleanup(self.formulas.close)

    def test_104_product_shares_match_the_reference_yaml_with_generic_carbonate(self):
        from primary_vs_secondary_metals.metals_primary_vs_secondary import read_excel_shares
        workbook_shares = read_excel_shares(self.path)
        yaml_shares = load_yaml(self.path.with_suffix(".yaml"))
        self.assertEqual(len(workbook_shares), 104)
        self.assertEqual(set(workbook_shares), set(yaml_shares))
        for product, entry in workbook_shares.items():
            with self.subTest(product=product):
                self.assertEqual(entry["name"], yaml_shares[product]["name"])
                self.assertEqual(entry["reference product"], yaml_shares[product]["reference product"])
                for role in ("primary", "secondary"):
                    self.assertEqual(set(entry["shares"][role]), {2020, 2050})
                    for year in (2020, 2050):
                        self.assertAlmostEqual(entry["shares"][role][year],
                                               yaml_shares[product]["shares"][role][year])

    def test_documented_eleven_market_assumptions_match_all_live_policy_cases(self):
        rows = list(self.values["Scenario assumptions"].iter_rows(
            min_row=6, max_row=16, max_col=7, values_only=True))
        self.assertEqual(len(rows), 11)
        self.assertEqual(len({row[0] for row in rows}), 11)
        policies = {case: MetalSteelPolicy.from_file(POLICY, nonferrous_case=case)
                    for case in ("reference", *CASE_INCREMENT)}
        rules_by_case = {case: {rule["reference product"]: rule for rule in selected.rules}
                         for case, selected in policies.items()}
        self.assertEqual({row[0] for row in rows}, set(rules_by_case["reference"]))
        for product, reference, anchor, low, middle, high, scope in rows:
            with self.subTest(product=product):
                rule = rules_by_case["reference"][product]
                self.assertAlmostEqual(reference, rule["reference_shares"])
                self.assertEqual({value.strip() for value in scope.split("/")}, set(rule["locations"]))
                if product in OBSERVED_2025:
                    self.assertAlmostEqual(anchor, OBSERVED_2025[product])
                else:
                    self.assertIsNone(anchor)
                for case, documented in zip(CASE_INCREMENT, (low, middle, high)):
                    curve = rules_by_case[case][product]["targets"][case]
                    self.assertAlmostEqual(documented, curve[2050])
                    if product in OBSERVED_2025:
                        self.assertAlmostEqual(anchor, curve[2025])

    def test_evidence_numbers_formula_text_and_cached_ratios_reconcile_independently(self):
        expected = {
            "gold": (3671.6, 1404.3, "wgc_gold_2025"),
            "silver": (846.6, 197.6, "silver_survey_2026"),
            "platinum": (5557, 1547, "jm_pgm_2026"),
            "palladium": (6589, 3135, "jm_pgm_2026"),
        }
        seen = set()
        for row in range(6, 10):
            product, year, primary, secondary, unit, ratio, source_id = next(
                self.values["Evidence"].iter_rows(min_row=row, max_row=row, max_col=7, values_only=True))
            with self.subTest(product=product):
                self.assertIn(product, expected)
                self.assertNotIn(product, seen)
                seen.add(product)
                self.assertEqual(year, 2025)
                self.assertTrue(unit)
                p, s, source = expected[product]
                self.assertAlmostEqual(primary, p)
                self.assertAlmostEqual(secondary, s)
                self.assertEqual(source_id, source)
                self.assertIsInstance(ratio, (int, float))
                self.assertAlmostEqual(ratio, s / (p + s))
                self.assertAlmostEqual(ratio, OBSERVED_2025[product])
                self.assertEqual(self.formulas["Evidence"][f"F{row}"].value,
                                 f"=D{row}/SUM(C{row}:D{row})")
        self.assertEqual(seen, set(expected))


if __name__ == "__main__":
    unittest.main()
