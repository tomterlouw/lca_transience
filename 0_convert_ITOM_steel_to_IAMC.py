"""Convert European ITOM steel exports to a reviewed premise IAMC resource.

The raw workbook is read-only. EAF-DRI chemistry is allocated using reported
domestic iron-route production as an explicit screening assumption. HBI uses
reported material sources: explicit NG/H2 imports and an EU production mix for
intra-EU trade. No energy total invents a technology-specific hydrogen share.
"""
from __future__ import annotations

import argparse
from datetime import datetime, timezone
import hashlib
from io import BytesIO
import json
from pathlib import Path
import re
import sys
import time
import xml.etree.ElementTree as ET
from uuid import uuid4
from zipfile import ZipFile, ZIP_DEFLATED, ZipInfo

import numpy as np
import pandas as pd
import pycountry

ROOT = Path(__file__).resolve().parent
DEFAULT_INPUT = ROOT / "inventories" / "Copy of ITOM_Steel_Test-scenario_IAMC_v2.xlsx"
DEFAULT_OUTPUT = ROOT / "scenario_data" / "scenario_data_itom.csv"
META = ["model", "scenario", "region", "variables", "unit"]
STEEL_PREFIX = "Production|Industry|Iron and Steel|Steel|"
IRON_PREFIX = "Production|Industry|Iron and Steel|Iron|"
MATERIAL_PREFIX = "Material Consumption|Industry|Iron and Steel|"
# Fixed EU27 modelling boundary for the intra-EU supplier pool. The UK is not
# included merely because it is present in the European ITOM workbook.
EU27 = frozenset("AT BE BG HR CY CZ DK EE FI FR DE GR HU IE IT LV LT LU MT NL PL PT RO SK SI ES SE".split())
# EU27 plus other whole-country locations in the installed RER geography.
# Cross-boundary countries need an explicit coverage decision before inclusion.
EUROPEAN_COUNTRIES = EU27 | frozenset("AD AL AX BA BY CH GB GG GI IM IS JE LI MC MD ME MK NO RS SJ SM UA VA".split())
EUROPEAN_REGION = "RER"
COUNTRY_SUPPLY_PREFIX = "Production|Steel|Supply from|"
HBI_SOURCES = ("HBI_intra_EU", "HBI_import_NG", "HBI_import_H2")
STEEL_ROUTES = ("BOF", "EAF_DRI", "EAF_HBI", "EAF_Scrap")
IRON_ROUTES = ("Blastfurnace", "DRI_NG", "DRI_H2", "DRI_NG_H2")
LEAVES = ("BF/BOF", "Secondary steel", "DRI_NG", "DRI_H2", "DRI_NG_H2",
          "DRI_unresolved", "HBI_NG", "HBI_H2", "HBI_NG_H2", "HBI_unresolved")
SHARE_COLUMNS = ("ng_share", "h2_share", "mixed_share", "unresolved_share")
REGION_ALIASES = {
    "slovokia": "SK", "slovakia": "SK", "slovak republic": "SK",
    "united kingdom": "GB", "uk": "GB", "great britain": "GB",
    "czech republic": "CZ", "czechia": "CZ", "greece": "GR", "el": "GR",
    "turkey": "TR", "türkiye": "TR", "south korea": "KR",
}


def normalize_region(value):
    """Resolve a country with exact ISO/name lookup; no fuzzy geographic guesses."""
    if pd.isna(value) or not str(value).strip():
        raise ValueError("Missing steel-export country.")
    raw = re.sub(r"\s+", " ", str(value).strip())
    if raw.casefold() in REGION_ALIASES:
        return REGION_ALIASES[raw.casefold()]
    try:
        return pycountry.countries.lookup(raw).alpha_2
    except LookupError as error:
        raise ValueError(f"Unsupported country {raw!r}; add an explicit, reviewed alias.") from error


def normalize_variable(value):
    if pd.isna(value):
        raise ValueError("Missing IAMC variable.")
    parts = [re.sub(r"\s+", " ", part.strip()) for part in str(value).split("|")]
    if any(not part for part in parts):
        raise ValueError(f"Empty IAMC hierarchy component in {value!r}.")
    return "|".join(parts)


def _years(frame):
    years = [str(c) for c in frame.columns if re.fullmatch(r"\d{4}", str(c))]
    if not years:
        raise ValueError("Steel IAMC input has no four-digit year columns.")
    if len(years) != len(set(years)):
        raise ValueError("Duplicate IAMC year columns.")
    return sorted(years, key=int)


def _validate_wide(frame):
    """Validate independently so the derivation API is safe without file reading."""
    out = frame.copy()
    out.columns = [str(c).strip() for c in out.columns]
    missing = set(META) - set(out)
    if missing:
        raise ValueError(f"IAMC input is missing columns: {sorted(missing)}")
    years = _years(out)
    extra = set(out) - set(META) - set(years)
    if extra:
        raise ValueError(f"Unexpected IAMC columns: {sorted(extra)}")
    for key in META:
        if out[key].isna().any() or out[key].astype(str).str.strip().eq("").any():
            raise ValueError(f"Missing {key} metadata in steel IAMC input.")
        out[key] = out[key].astype(str).str.strip()
    if out.duplicated(META).any():
        raise ValueError("Duplicate IAMC identities; refusing to sum or discard repeated input rows.")
    for year in years:
        values = pd.to_numeric(out[year], errors="raise").astype(float)
        if not np.isfinite(values.to_numpy()).all():
            raise ValueError(f"Missing or non-finite values in {year}; blank year cells are not zero.")
        if values.lt(0).any():
            raise ValueError(f"Negative production/energy/material amount in {year}.")
        out[year] = values
    return out[META + years]


def read_itom_steel_workbook(path, *, model_case="lower"):
    """Read all three raw sheets and return normalized IAMC plus country audit."""
    if model_case not in ("lower", "preserve"):
        raise ValueError("model_case must be 'lower' or 'preserve'.")
    frames, country_rows = [], []
    with pd.ExcelFile(Path(path)) as book:
        required = {"Production", "Final Energy", "Material Consumption"}
        if required - set(book.sheet_names):
            raise ValueError(f"Workbook lacks sheets: {sorted(required-set(book.sheet_names))}")
        for sheet in ("Production", "Final Energy", "Material Consumption"):
            raw = pd.read_excel(book, sheet_name=sheet).dropna(how="all")
            columns = [str(c).strip().lower() for c in raw.columns]
            if len(columns) != len(set(columns)):
                raise ValueError(f"Ambiguous duplicate headers in {sheet}.")
            raw.columns = columns
            variable_names = [name for name in ("name", "variable", "variables") if name in raw]
            if len(variable_names) != 1:
                raise ValueError(f"{sheet} needs exactly one Name/Variable/Variables column.")
            raw = raw.rename(columns={variable_names[0]: "variables"})
            missing = set(META) - set(raw)
            if missing:
                raise ValueError(f"{sheet} is missing columns: {sorted(missing)}")
            if raw.model.isna().any():
                raise ValueError(f"{sheet} has missing model metadata.")
            for original in raw["region"].drop_duplicates().tolist():
                country_rows.append(dict(original_region=str(original), iso2=normalize_region(original)))
            raw["region"] = raw.region.map(normalize_region)
            raw["variables"] = raw.variables.map(normalize_variable)
            if model_case == "lower": raw["model"] = raw.model.astype(str).str.strip().str.lower()
            expected_unit = "PJ/yr" if sheet == "Final Energy" else "Mt/yr"
            units = raw.unit.astype(str).str.strip()
            if not units.eq(expected_unit).all():
                raise ValueError(f"{sheet} requires explicit {expected_unit} units; found {sorted(units.unique())}.")
            raw["unit"] = expected_unit
            frames.append(_validate_wide(raw))
    result = _validate_wide(pd.concat(frames, ignore_index=True))
    return result, pd.DataFrame(country_rows).drop_duplicates().sort_values("original_region").reset_index(drop=True)


def _override_table(overrides):
    if overrides is None:
        return pd.DataFrame(columns=["scenario", "region", "year", "route"] + list(SHARE_COLUMNS))
    out = overrides.copy()
    needed = {"scenario", "region", "year", "route"} | set(SHARE_COLUMNS)
    if needed - set(out):
        raise ValueError(f"Route overrides are missing columns: {sorted(needed-set(out))}")
    for key in ("scenario", "region", "year", "route"):
        if out[key].isna().any(): raise ValueError(f"Missing override {key}; use '*' explicitly for a wildcard.")
        out[key] = out[key].astype(str).str.strip()
    out["region"] = out.region.map(lambda x: x if x == "*" else normalize_region(x))
    if not set(out.route).issubset({"EAF_DRI", "EAF_HBI"}):
        raise ValueError("Override route must be EAF_DRI or EAF_HBI.")
    for year in out.year:
        if year != "*" and (not re.fullmatch(r"\d{4}(?:\.0)?", year)):
            raise ValueError(f"Invalid override year {year!r}.")
    out["year"] = out.year.map(lambda x: x if x == "*" else str(int(float(x))))
    for key in SHARE_COLUMNS:
        values = pd.to_numeric(out[key], errors="raise").astype(float)
        if not np.isfinite(values).all() or not values.between(0, 1).all():
            raise ValueError(f"Override {key} values must be finite fractions from zero to one.")
        out[key] = values
    if not np.allclose(out[list(SHARE_COLUMNS)].sum(axis=1), 1., rtol=0, atol=1e-9):
        raise ValueError("Override shares must sum to one.")
    if out.duplicated(["scenario", "region", "year", "route"]).any():
        raise ValueError("Duplicate override identities.")
    return out.reset_index(drop=True)


def _select_override(overrides, scenario, region, year, route):
    if overrides.empty: return None, None
    chosen = overrides.loc[overrides.route.eq(route)].copy()
    chosen["_specificity"] = 0
    for key, target in (("scenario", scenario), ("region", region), ("year", str(year))):
        matches = chosen[key].eq(str(target))
        chosen = chosen.loc[matches | chosen[key].eq("*")].copy()
        chosen["_specificity"] += chosen[key].ne("*").astype(int)
    if chosen.empty: return None, None
    chosen = chosen.loc[chosen._specificity.eq(chosen._specificity.max())]
    if len(chosen) != 1:
        raise ValueError(f"Ambiguous route overrides for {scenario}/{region}/{year}/{route}.")
    return chosen.iloc[0][list(SHARE_COLUMNS)].to_numpy(dtype=float), int(chosen.index[0])


def _hbi_source_context(raw):
    """Build material-source and EU iron-production lookups, without mixing cases."""
    materials = raw.loc[raw.variables.str.startswith(MATERIAL_PREFIX + "HBI")].copy()
    unknown = set(materials.variables) - {MATERIAL_PREFIX + source for source in HBI_SOURCES}
    if unknown:
        raise ValueError(f"Unsupported HBI material source(s): {sorted(unknown)}; define their provenance before conversion.")
    if not materials.unit.eq("Mt/yr").all():
        raise ValueError("HBI material consumption must be Mt/yr.")
    sources = {identity: part.set_index("variables") for identity, part in
               materials.groupby(["model", "scenario", "region"], sort=True)}
    eu_iron = raw.loc[raw.region.isin(EU27) & raw.variables.isin(
        [IRON_PREFIX + route for route in ("DRI_NG", "DRI_H2", "DRI_NG_H2")])]
    if not eu_iron.unit.eq("Mt/yr").all():
        raise ValueError("EU DRI iron production must be Mt/yr.")
    pools = {identity: part.groupby("variables")[_years(raw)].sum() for identity, part in
             eu_iron.groupby(["model", "scenario"], sort=True)}
    return sources, pools


def _hbi_fractions(sources, pools, identity, year):
    """Allocate existing EAF-HBI steel, using material shares, not material tonnes."""
    material = sources.get(identity)
    amounts = np.array([
        float(material.loc[MATERIAL_PREFIX + source, year])
        if material is not None and MATERIAL_PREFIX + source in material.index else 0.
        for source in HBI_SOURCES
    ])
    pool = pools.get(identity[:2])
    eu = np.array([
        float(pool.loc[IRON_PREFIX + route, year])
        if pool is not None and IRON_PREFIX + route in pool.index else 0.
        for route in ("DRI_NG", "DRI_H2", "DRI_NG_H2")
    ])
    details = dict(zip((source + "_Mt_per_year" for source in HBI_SOURCES), map(float, amounts)))
    details.update(dict(zip(("EU27_DRI_NG_Mt_per_year", "EU27_DRI_H2_Mt_per_year",
                             "EU27_DRI_NG_H2_Mt_per_year"), map(float, eu))))
    details.update(HBI_material_total_Mt_per_year=float(amounts.sum()),
                   EU27_DRI_total_Mt_per_year=float(eu.sum()),
                   HBI_material_rows_present="; ".join(
                       source for source in HBI_SOURCES
                       if material is not None and MATERIAL_PREFIX + source in material.index))
    if amounts.sum() == 0:
        return np.array([0., 0., 0., 1.]), "HBI_no_positive_material_source_unresolved", details
    # Explicit extra-EU sources retain their reported reducing agent. Intra-EU
    # HBI receives the contemporaneous EU27 DRI iron mix, a declared proxy for
    # unknown exporters, with mixed-route chemistry still explicitly separate.
    allocated = np.array([amounts[1], amounts[2], 0., 0.])
    if amounts[0] > 0 and eu.sum() > 0:
        allocated[:3] += amounts[0] * eu / eu.sum()
        method = "material_hbi_sources_with_eu_dri_mix_assumption"
    elif amounts[0] > 0:
        allocated[3] = amounts[0]
        method = "material_hbi_sources_missing_eu_dri_mix_unresolved"
    else:
        method = "material_hbi_sources_explicit_imports"
    return allocated / amounts.sum(), method, details


def derive_steel_routes(raw, overrides=None, *, strict_missing_routes=False):
    """Derive mutually exclusive steel leaves while conserving source steel.

    No iron volume is relabelled as steel. A country's reported EAF_DRI steel
    volume is split using its domestic iron-route production fractions. This
    proportional-routing assumption cannot resolve trade or a plant's actual
    fuel mix. HBI material sources determine route shares; intra-EU HBI uses
    the modelled EU27 DRI mix as a documented assumption, not traced trade.
    """
    raw = _validate_wide(raw)
    raw['region'] = raw.region.map(normalize_region)
    raw['variables'] = raw.variables.map(normalize_variable)
    raw = _validate_wide(raw)
    years = _years(raw)
    production = raw.loc[raw.variables.str.startswith(STEEL_PREFIX) | raw.variables.str.startswith(IRON_PREFIX)].copy()
    if production.empty: raise ValueError("No ITOM iron/steel production variables found.")
    if not production.unit.eq("Mt/yr").all(): raise ValueError("Iron and steel production must be Mt/yr.")
    for variable in production.variables:
        if variable.startswith(STEEL_PREFIX) and variable.removeprefix(STEEL_PREFIX) not in STEEL_ROUTES:
            raise ValueError(f"Unsupported steel route {variable!r}; define its allocation before conversion.")
        if variable.startswith(IRON_PREFIX) and variable.removeprefix(IRON_PREFIX) not in IRON_ROUTES:
            raise ValueError(f"Unsupported iron route {variable!r}.")
    override_rules = _override_table(overrides)
    hbi_sources, eu_pools = _hbi_source_context(raw)
    used_overrides, records, audit = set(), [], []
    for identity, part in production.groupby(["model", "scenario", "region"], sort=True):
        model, scenario, region = identity
        values = part.set_index("variables")
        missing = [prefix + route for prefix, routes in ((STEEL_PREFIX, STEEL_ROUTES), (IRON_PREFIX, IRON_ROUTES))
                   for route in routes if prefix + route not in values.index]
        if missing and strict_missing_routes:
            raise ValueError(f"Missing whole production routes in {scenario}/{region}: {missing}")
        outputs = {leaf: {"model": model, "scenario": scenario, "region": region,
                          "variables": "Production|Steel|" + leaf, "unit": "Mt/yr"} for leaf in LEAVES}
        total_row = dict(model=model, scenario=scenario, region=region, variables="Production|Steel", unit="Mt/yr")
        for year in years:
            def quantity(prefix, route):
                name = prefix + route
                return float(values.loc[name, year]) if name in values.index else 0.
            steel = {route: quantity(STEEL_PREFIX, route) for route in STEEL_ROUTES}
            iron = {route: quantity(IRON_PREFIX, route) for route in IRON_ROUTES}
            total = sum(steel.values())
            outputs["BF/BOF"][year] = steel["BOF"]
            outputs["Secondary steel"][year] = steel["EAF_Scrap"]
            domestic = np.array([iron["DRI_NG"], iron["DRI_H2"], iron["DRI_NG_H2"]], dtype=float)
            for route, prefix in (("EAF_DRI", "DRI"), ("EAF_HBI", "HBI")):
                hbi_details = {}
                if route == "EAF_HBI":
                    hbi_fractions, hbi_method, hbi_details = _hbi_fractions(hbi_sources, eu_pools, identity, year)
                fractions, override_index = _select_override(override_rules, scenario, region, year, route)
                if fractions is not None:
                    method = "explicit_override"; used_overrides.add(override_index)
                elif route == "EAF_DRI" and domestic.sum() > 0:
                    fractions = np.r_[domestic / domestic.sum(), 0.]
                    method = "domestic_iron_proportional_allocation_assumption"
                elif route == "EAF_HBI":
                    fractions, method = hbi_fractions, hbi_method
                else:
                    fractions = np.array([0., 0., 0., 1.])
                    method = "no_domestic_iron_origin_unresolved"
                if not np.isclose(fractions.sum(), 1., atol=1e-9, rtol=0) or (fractions < 0).any():
                    raise ValueError(f"Invalid derived {route} fractions in {scenario}/{region}/{year}.")
                amounts = float(steel[route]) * fractions
                for leaf, amount in zip((prefix + "_NG", prefix + "_H2", prefix + "_NG_H2", prefix + "_unresolved"), amounts):
                    outputs[leaf][year] = float(amount)
                audit.append(dict(model=model, scenario=scenario, region=region, year=int(year), route=route,
                    source_steel_Mt_per_year=steel[route], domestic_DRI_NG_Mt_per_year=float(domestic[0]),
                    domestic_DRI_H2_Mt_per_year=float(domestic[1]), domestic_DRI_NG_H2_Mt_per_year=float(domestic[2]),
                    domestic_iron_total_Mt_per_year=float(domestic.sum()), method=method,
                    **dict(zip(SHARE_COLUMNS, map(float, fractions))),
                    derived_steel_Mt_per_year=float(amounts.sum()),
                    **hbi_details,
                    missing_whole_rows_assumed_zero="; ".join(missing),
                    iron_vs_EAF_ratio=(None if route != "EAF_DRI" or steel[route] == 0 else
                                       float(domestic.sum() / steel[route])),
                    HBI_material_to_steel_ratio=(None if route != "EAF_HBI" or steel[route] == 0 else
                                                hbi_details["HBI_material_total_Mt_per_year"] / steel[route]),
                    warning=("Domestic iron/steel production is not a closed trade/yield balance; routing is inferred. "
                             "Intra-EU HBI uses an EU27 production-mix assumption, not traced exporters. "
                             "HBI material quantities determine shares, not additional steel output. "
                             "The hybrid NG/H2 fuel fraction remains unresolved. No CCS/electrolysis volume is invented.")))
            derived_total = sum(outputs[leaf][year] for leaf in LEAVES)
            if not np.isclose(derived_total, total, atol=1e-9, rtol=1e-10):
                raise ValueError(f"Derived steel route total does not reconcile for {scenario}/{region}/{year}.")
            total_row[year] = total
        records.append(total_row); records.extend(outputs.values())
    unused = set(override_rules.index) - used_overrides
    if unused: raise ValueError(f"Override rows match no requested country/year/route: {sorted(unused)}")
    derived = _validate_wide(pd.DataFrame(records))
    return derived, pd.DataFrame(audit)


def aggregate_european_steel(derived):
    """Sum national steel routes and expose their volume-weighted contributions.

    Only the eleven canonical country-level steel rows are accepted. RER and
    country rows must never be added together as independent production pools.
    """
    countries = _validate_wide(derived)
    expected = {"Production|Steel"} | {"Production|Steel|" + leaf for leaf in LEAVES}
    if not set(countries.variables).issubset(expected):
        raise ValueError("European aggregation accepts only canonical steel rows, not raw iron/material/energy rows.")
    if not countries.unit.eq("Mt/yr").all():
        raise ValueError("European steel production must be Mt/yr.")
    unsupported = set(countries.region) - EUROPEAN_COUNTRIES
    if unsupported:
        raise ValueError(f"European aggregation requires reviewed European ISO2 countries, not aggregates or other regions: {sorted(unsupported)}")
    years = _years(countries)
    aggregates, shares = [], []
    for (model, scenario), context in countries.groupby(["model", "scenario"], sort=True):
        for region, part in context.groupby("region", sort=True):
            if set(part.variables) != expected:
                raise ValueError(f"Incomplete canonical steel routes for {model}/{scenario}/{region}.")
            amounts = part.set_index("variables")[years]
            total = amounts.loc["Production|Steel"]
            leaf_sum = amounts.drop(index="Production|Steel").sum()
            if not np.allclose(total, leaf_sum, atol=1e-9, rtol=1e-10):
                raise ValueError(f"Country steel total does not reconcile in {model}/{scenario}/{region}.")
        summed = context.groupby("variables", sort=True)[years].sum()
        for variable, amounts in summed.iterrows():
            aggregates.append(dict(model=model, scenario=scenario, region=EUROPEAN_REGION,
                variables=variable, unit="Mt/yr", **amounts.to_dict()))
        for year in years:
            european_total = float(summed.loc["Production|Steel", year])
            for region, part in context.groupby("region", sort=True):
                amounts = part.set_index("variables")
                country_total = float(amounts.loc["Production|Steel", year])
                for route in LEAVES:
                    variable = "Production|Steel|" + route
                    volume = float(amounts.loc[variable, year])
                    route_volume = float(summed.loc[variable, year])
                    shares.append(dict(model=model, scenario=scenario, year=int(year),
                        aggregate_region=EUROPEAN_REGION, region=region, route=route,
                        variables=variable, unit="Mt/yr", production_Mt_per_year=volume,
                        country_steel_Mt_per_year=country_total,
                        european_route_Mt_per_year=route_volume,
                        european_steel_Mt_per_year=european_total,
                        country_share=country_total / european_total if european_total > 0 else np.nan,
                        route_share=route_volume / european_total if european_total > 0 else np.nan,
                        country_route_share=volume / european_total if european_total > 0 else np.nan,
                        country_share_within_route=volume / route_volume if route_volume > 0 else np.nan,
                        status=("no_european_production" if european_total == 0 else
                                "no_route_production" if route_volume == 0 else "defined")))
    if not aggregates:
        raise ValueError("No country steel production available for European aggregation.")
    return _validate_wide(pd.DataFrame(aggregates)), pd.DataFrame(shares)


def european_country_supply_iamc(derived):
    """Export each country's total as an RER national-market supply variable."""
    # Reuse the full conservation/coverage checks before exposing a second view.
    aggregate_european_steel(derived)
    totals = _validate_wide(derived)
    totals = totals.loc[totals.variables.eq("Production|Steel")].copy()
    totals["variables"] = COUNTRY_SUPPLY_PREFIX + totals.region
    totals["region"] = EUROPEAN_REGION
    return _validate_wide(totals).sort_values(["model", "scenario", "variables"], kind="stable").reset_index(drop=True)


def _iamc_excel_bytes(frame, shares):
    """Export the same IAMC rows and the transparent production-weight audit.

    This uses the repository's ordinary Excel engine so the converter remains a
    self-contained Python command. Stable ZIP metadata permits unchanged reruns
    while a user has the existing workbook open in Excel.
    """
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    country_mix = shares.drop_duplicates(["model", "scenario", "year", "region"])[[
        "model", "scenario", "year", "aggregate_region", "region", "country_steel_Mt_per_year",
        "european_steel_Mt_per_year", "country_share"]].copy()
    country_mix = country_mix.rename(columns={"region": "source_country"})
    country_mix["low_alloyed_supplier"] = "market for steel, low-alloyed (ITOM)"
    country_mix["unalloyed_supplier"] = "market for steel, unalloyed (ITOM)"
    stream = BytesIO()
    with pd.ExcelWriter(stream, engine="openpyxl") as writer:
        frame.to_excel(writer, sheet_name="IAMC", index=False)
        country_mix.to_excel(writer, sheet_name="European country mix", index=False, na_rep="")
        shares.to_excel(writer, sheet_name="European market shares", index=False, na_rep="")
        for name, table in (("IAMC", frame), ("European country mix", country_mix), ("European market shares", shares)):
            sheet = writer.sheets[name]
            sheet.freeze_panes = "F2" if name in ("IAMC", "European country mix") else "I2"
            sheet.auto_filter.ref = sheet.dimensions
            sheet.row_dimensions[1].height = 40
            for column, field in enumerate(table.columns, 1):
                cell = sheet.cell(1, column)
                cell.font = Font(name="Calibri", bold=True, color="FFFFFF")
                cell.fill = PatternFill("solid", fgColor="274454")
                cell.alignment = Alignment(wrap_text=True, vertical="center")
                width = (68 if field == "variables" or str(field).endswith("_supplier") else 34 if field == "scenario" else
                         24 if field in ("model", "route", "status") else
                         23 if "share" in str(field) or "Mt_per_year" in str(field) else 15)
                if name != "IAMC":
                    width = max(width, len(str(field)) + 2)
                sheet.column_dimensions[get_column_letter(column)].width = width
                for row in range(2, len(table) + 2):
                    value_cell = sheet.cell(row, column)
                    if "share" in str(field):
                        value_cell.number_format = "0.00%"
                    elif str(field).isdigit() or "Mt_per_year" in str(field):
                        value_cell.number_format = "0.000000"
    # Avoid timestamps making a data-identical workbook appear changed. Build
    # timestamps and source hashes remain in the separate conversion manifest.
    stable = BytesIO()
    with ZipFile(BytesIO(stream.getvalue())) as source, ZipFile(stable, "w", ZIP_DEFLATED) as target:
        for entry in source.infolist():
            data = source.read(entry.filename)
            if entry.filename == "docProps/core.xml":
                core = ET.fromstring(data)
                for node in list(core):
                    if node.tag.rsplit("}", 1)[-1] in ("created", "modified"):
                        core.remove(node)
                data = ET.tostring(core, encoding="utf-8")
            normalized = ZipInfo(entry.filename, date_time=(1980, 1, 1, 0, 0, 0))
            normalized.compress_type = ZIP_DEFLATED
            normalized.external_attr = entry.external_attr
            target.writestr(normalized, data)
    return stable.getvalue()


def _sha256(path):
    digest = hashlib.sha256()
    with Path(path).open("rb") as stream:
        for block in iter(lambda: stream.read(1024 * 1024), b""): digest.update(block)
    return digest.hexdigest()


def transform_itom_steel(input_path=DEFAULT_INPUT, output_path=DEFAULT_OUTPUT, *,
                         overrides_path=None, strict_missing_routes=False, model_case="lower"):
    """Publish country/RER IAMC as CSV and Excel, with route/share audits."""
    source, destination = Path(input_path).expanduser().resolve(), Path(output_path).expanduser().resolve()
    if source == destination: raise ValueError("Output cannot replace the raw workbook.")
    if destination.suffix.lower() not in (".csv", ".xlsx"):
        raise ValueError("IAMC output must use a .csv or .xlsx extension.")
    destination = destination.with_suffix(".csv")
    xlsx_path = destination.with_suffix(".xlsx")
    if source == xlsx_path: raise ValueError("Excel output cannot replace the raw workbook.")
    source_hash = _sha256(source)
    raw, countries = read_itom_steel_workbook(source, model_case=model_case)
    overrides_file = None if overrides_path is None else Path(overrides_path).expanduser().resolve()
    if overrides_file is not None and overrides_file == destination:
        raise ValueError("Output cannot replace the route overrides input.")
    override_hash = None if overrides_file is None else _sha256(overrides_file)
    overrides = None if overrides_file is None else pd.read_csv(overrides_file, dtype={"year": str})
    derived, routes = derive_steel_routes(raw, overrides, strict_missing_routes=strict_missing_routes)
    european, shares = aggregate_european_steel(derived)
    country_supplies = european_country_supply_iamc(derived)
    combined = _validate_wide(pd.concat([raw, derived, european, country_supplies], ignore_index=True))
    combined = combined.sort_values(["model", "scenario", "region", "variables"], kind="stable").reset_index(drop=True)
    if _sha256(source) != source_hash or (overrides_file is not None and _sha256(overrides_file) != override_hash):
        raise RuntimeError("An input changed while converting; no output published.")
    csv_bytes = combined.to_csv(index=False).encode("utf-8-sig")
    audit_bytes = routes.to_csv(index=False).encode("utf-8-sig")
    shares_bytes = shares.to_csv(index=False).encode("utf-8-sig")
    excel_bytes = _iamc_excel_bytes(combined, shares)
    audit_path = destination.with_name(destination.stem + "_route_audit.csv")
    shares_path = destination.with_name(destination.stem + "_market_shares.csv")
    manifest_path = destination.with_name(destination.stem + "_conversion_manifest.json")
    sidecars = {audit_path, shares_path, xlsx_path, manifest_path}
    if source in sidecars or (overrides_file is not None and overrides_file in sidecars):
        raise ValueError("An output sidecar would replace an input.")
    manifest = dict(created_utc=datetime.now(timezone.utc).isoformat(), source=str(source), source_sha256=source_hash,
        output=str(destination), output_sha256=hashlib.sha256(csv_bytes).hexdigest(),
        audit=str(audit_path), audit_sha256=hashlib.sha256(audit_bytes).hexdigest(),
        xlsx=str(xlsx_path), xlsx_sha256=hashlib.sha256(excel_bytes).hexdigest(),
        market_shares=str(shares_path), market_shares_sha256=hashlib.sha256(shares_bytes).hexdigest(),
        overrides=None if overrides_file is None else str(overrides_file), overrides_sha256=override_hash,
        raw_rows=int(len(raw)), derived_rows=int(len(derived)), european_rows=int(len(european) + len(country_supplies)),
        european_route_rows=int(len(european)), european_country_supply_rows=int(len(country_supplies)), output_rows=int(len(combined)),
        country_count=int(raw.region.nunique()), countries=countries.to_dict("records"),
        years=list(map(int, _years(combined))), models=combined.model.unique().tolist(),
        scenarios=combined.scenario.unique().tolist(), units=combined.unit.unique().tolist(),
        raw_sheets=["Production", "Final Energy", "Material Consumption"],
        european_aggregate=dict(region=EUROPEAN_REGION, country_only_statistics=True,
            contexts=[dict(model=model, scenario=scenario, countries=sorted(part.region.unique()))
                      for (model, scenario), part in derived.groupby(["model", "scenario"], sort=True)],
            route_rows=european.to_dict("records"),
            weighting="national total steel production divided by total European steel production",
            country_supply_prefix=COUNTRY_SUPPLY_PREFIX,
            country_supply_rows=country_supplies.to_dict("records"),
            market_inventory_geography="national ITOM market supply; direct same-grade country-market inputs",
            post_update_step_required="itom_steel_country_market_mix in existing export workflow",
            excluded_from_pooling=["HBI EU27 origin inference", "country counts", "country-only route totals"],
            automatic_world_market="excluded to avoid adding RER and countries as independent production"),
        hbi_source_contract=dict(
            material_variables=[MATERIAL_PREFIX + source for source in HBI_SOURCES],
            intra_eu_mix="production-weighted EU27 DRI iron, same model/scenario/year",
            modelled_eu27_countries=sorted(set(raw.region) & EU27),
            boundary_eu27_countries=sorted(EU27),
            unavailable_eu27_countries=sorted(EU27 - set(raw.region)),
            material_basis="relative source quantities allocate reported EAF_HBI steel; no assumed 1:1 yield",
            origin_geography="exporters not traced; integrated HBI LCIs retain geography/briquetting/transport proxies"),
        method_counts=routes.method.value_counts().to_dict(),
        positive_route_diagnostics={
            "HBI_unresolved_country_year_cells": int(((routes.route == 'EAF_HBI') & routes.source_steel_Mt_per_year.gt(0) & routes.unresolved_share.gt(0)).sum()),
            "DRI_unresolved_country_year_cells": int(((routes.route == 'EAF_DRI') & routes.source_steel_Mt_per_year.gt(0) & routes.unresolved_share.gt(0)).sum()),
            "mixed_route_country_year_cells": int((routes.source_steel_Mt_per_year.gt(0) & routes.mixed_share.gt(0)).sum()),
        },
        derived_route_totals_Mt_per_year={str(variable): {year: float(part[year].sum()) for year in _years(derived)}
                                        for variable, part in derived.groupby('variables', sort=True)},
        missing_whole_route_policy="fail" if strict_missing_routes else "zero with route-audit evidence",
        assumptions=[
            "Raw Mt/yr and PJ/yr values are already annual; no five-year division or magnitude conversion is applied.",
            "Only column headers and default model case are lowercased. ISO2 codes and canonical IAMC hierarchy case are preserved.",
            "EAF_DRI steel volume is split in proportion to domestic DRI_NG/DRI_H2/DRI_NG_H2 iron production unless explicitly overridden.",
            "That split is a routing assumption, not observed feedstock use, trade closure or an iron-to-steel yield.",
            "HBI_import_NG/H2 material quantities identify extra-EU reduction routes; HBI_intra_EU receives the same model/scenario/year EU27 DRI production mix. Exporters are not traced.",
            "Only modelled EU27 countries contribute to that mix; missing countries are not filled from other scenarios or national energy totals. The UK is outside this pool.",
            "HBI steel volumes are divided by relative material-source amounts, conserving reported EAF_HBI steel. Material input tonnes are not renamed as steel or added as another production stream.",
            "Intra-EU HBI without a positive EU DRI denominator, or HBI steel without positive material-source data, remains unresolved. Explicit reviewed route overrides take precedence.",
            "Mixed NG/H2 output is kept separate; its hydrogen/gas input fractions are not inferred.",
            "BF/BOF and full-scrap EAF route volumes are retained as high-level production shares, not finished-steel recycled-content fractions. ITOM's downstream dummy blend is not added again.",
            "Original normalized production, final-energy and material-consumption rows are preserved as diagnostics; do not add iron to steel output.",
            "RER is the sum of canonical steel-route production over the supplied European countries, including GB. It is not a separate source of additional production and does not fill missing countries.",
            "European route shares use production volumes, not an equal-country average. Raw iron, energy and intra-EU material-trade rows are not added to aggregate steel output.",
            "Country and RER rows coexist as alternative geographic views; retain World exclusion in the market mapping. Country-only totals and HBI-origin pooling exclude RER.",
            "RER country-supply variables allocate existing European steel production to national ITOM markets; these are not additional steel tonnes.",
            "The existing export post-step replaces provisional RER technology inputs with direct same-grade national ITOM markets, preserving national technology and energy mixes.",
            "Intermediate-year country production volumes are interpolated before normalization. Shares themselves are not interpolated.",
            "CCS is in ITOM but no route-specific capture data are available here; the current mapping does not represent its deployment. Electrowinning is available in ITOM but absent from this test scenario.",
            "National energy/material totals are diagnostics only. Route LCIs are not recalibrated without a separately supplied InputActivityRatio table and output-basis/mode mapping.",
            "This script does not update a database or enable the ITOM module in notebook 1.",
        ])
    payloads = {destination: csv_bytes, audit_path: audit_bytes, shares_path: shares_bytes, xlsx_path: excel_bytes,
                manifest_path: (json.dumps(manifest, indent=2, ensure_ascii=False, allow_nan=False) + "\n").encode("utf-8")}
    destination.parent.mkdir(parents=True, exist_ok=True)
    temporary = []
    try:
        for target, payload in payloads.items():
            stage = target.with_name("." + target.name + "." + uuid4().hex + ".tmp")
            stage.write_bytes(payload); temporary.append((stage, target))
        for stage, target in temporary:
            # Preserve byte-identical resources, including a CSV open in Excel.
            # The canonical descriptor path must never silently change.
            if target.is_file() and target.read_bytes() == stage.read_bytes():
                stage.unlink()
                continue
            for attempt, delay in enumerate((0., .2, .5, 1.)):
                if delay: time.sleep(delay)
                try:
                    stage.replace(target)
                    break
                except PermissionError as error:
                    if attempt == 3:
                        raise PermissionError(
                            f"Cannot replace {target}. Close Excel/viewer handles or check write permissions, "
                            "then rerun the converter. The raw workbook was not changed."
                        ) from error
    finally:
        for stage, _ in temporary: stage.unlink(missing_ok=True)
    return dict(output=str(destination), xlsx=str(xlsx_path), audit=str(audit_path),
                market_shares=str(shares_path), manifest=str(manifest_path), raw_rows=len(raw),
                derived_rows=len(derived), european_rows=len(european) + len(country_supplies),
                european_route_rows=len(european), european_country_supply_rows=len(country_supplies),
                output_rows=len(combined), country_count=raw.region.nunique())


def main(argv=None):
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--input", type=Path, default=DEFAULT_INPUT, help="Raw ITOM steel workbook (read-only).")
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT, help="IAMC output stem (.csv or .xlsx); both formats are written.")
    parser.add_argument("--route-overrides", type=Path, help="Reviewed NG/H2/mixed/unresolved shares by route/country/year.")
    parser.add_argument("--strict-missing-routes", action="store_true", help="Reject sparse missing whole route rows instead of audited zero assumptions.")
    parser.add_argument("--preserve-model-case", action="store_true", help="Keep original model name capitalization; headers always become lowercase.")
    if argv is None and Path(sys.argv[0]).stem == "ipykernel_launcher":
        parser.add_argument("-f", "--f", dest="_kernel_connection", help=argparse.SUPPRESS)
    args = parser.parse_args(argv)
    result = transform_itom_steel(args.input, args.output, overrides_path=args.route_overrides,
        strict_missing_routes=args.strict_missing_routes, model_case="preserve" if args.preserve_model_case else "lower")
    print(f"Converted {result['raw_rows']} original + {result['derived_rows']} derived rows across {result['country_count']} ISO2 countries.")
    if result.get("european_rows"):
        print(f"Added {result['european_rows']} RER rows as a separate European production view.")
    print("IAMC CSV:", result['output']); print("Route audit:", result['audit']); print("Source/assumption manifest:", result['manifest'])
    if "xlsx" in result: print("IAMC Excel:", result['xlsx'])
    if "market_shares" in result: print("European production weights:", result['market_shares'])
    return result


if __name__ == "__main__":
    main()
